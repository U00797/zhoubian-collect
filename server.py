#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Local bridge: browser page -> result bundle -> optional Codex export."""

import argparse
import csv
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from html import unescape
from html.parser import HTMLParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from math import ceil, floor
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, quote, unquote, urljoin, urlparse
from urllib.request import Request, urlopen

ROOT = Path(__file__).resolve().parent
CONFIG_FILE = ROOT / "config.json"
TARGETS = {
    "obsidian": "Obsidian",
    "wps": "WPS",
    "evernote": "印象笔记",
    "other": "其他软件 / 文件夹",
}
MODES = {
    "ai": "AI 自动识别",
    "local": "无 AI 本地导出",
}
EXPORT_COLUMNS = (
    ("ip", "IP"),
    ("publisher", "出品方"),
    ("release_date", "发售时间及地点"),
    ("series", "系列"),
    ("items", "制品明细"),
    ("spec", "尺寸丨材质丨工艺"),
    ("price", "价格"),
    ("images", "图片"),
    ("notes", "备注"),
    ("source_text", "原文"),
)
WEB_COLUMNS = (
    ("publisher", "出品方"),
    ("release_date", "发售时间及地点"),
    ("series", "系列"),
    ("ip", "IP名称"),
    ("items", "周边明细"),
    ("spec", "尺寸丨材质丨工艺"),
    ("price", "价格"),
    ("images", "图片"),
    ("notes", "备注"),
)
DEFAULT_CONFIG = {
    "host": "127.0.0.1",
    "port": 8768,
    "codex_cli": os.environ.get("CODEX_CLI_PATH", ""),
    "codex_thread": os.environ.get("CODEX_THREAD_ID", ""),
    "output_dir": "output",
}
PRICE_RE = re.compile(
    r"(?:(?:¥|￥)\s*)?(\d+(?:\.\d{1,2})?)\s*元"
    r"(?:\s*[/／]\s*(个|对|套|组|枚|张|盒|款|份))?",
    re.IGNORECASE,
)
UNIT_SALE_RE = re.compile(
    r"^\s*(\d+\s*(?:只|个|件|枚|对|套|组|盒|包|款|份)"
    r"\s*[/／]\s*套[，,、；;]?\s*按套售卖)\s*[；;]?\s*"
)
IMAGE_BAD_RE = re.compile(r"data:|\.(?:svg|gif)(?:\?|$)", re.I)
IMAGE_AVATAR_RE = re.compile(r"avatar|profile_pic|orj360|thumb150|/(?:30|50)/", re.I)
VISION_SLICE_HEIGHT = 3000


def merge_defaults(cfg):
    for key, value in DEFAULT_CONFIG.items():
        cfg.setdefault(key, value)
    return cfg


def load_config():
    cfg = {}
    if CONFIG_FILE.exists():
        cfg = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    merge_defaults(cfg)
    if not CONFIG_FILE.exists():
        CONFIG_FILE.write_text(
            json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    return cfg


def resolve_codex_cli(cfg):
    candidates = [cfg.get("codex_cli"), os.environ.get("CODEX_CLI_PATH")]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    local = Path(os.environ.get("LOCALAPPDATA", "")) / "OpenAI" / "Codex" / "bin"
    exes = sorted(local.glob("*/codex.exe"), key=lambda p: p.stat().st_mtime, reverse=True)
    if exes:
        return str(exes[0])
    return shutil.which("codex.exe") or shutil.which("codex") or ""


def target_name(value):
    return TARGETS.get(str(value or "").lower(), str(value or "指定软件"))


def mode_name(value):
    return MODES.get(str(value or "").lower(), MODES["ai"])


def clean_lines(text):
    lines = []
    for raw in str(text or "").splitlines():
        line = " ".join(raw.replace("\u3000", " ").split()).strip()
        if line and line not in lines:
            lines.append(line)
    return lines


def first_match(lines, pattern):
    rx = re.compile(pattern, re.IGNORECASE)
    for line in lines:
        match = rx.search(line)
        if match:
            value = match.group(1) if match.groups() else match.group(0)
            return value.strip(" ：:；;，,。")
    return ""


def keyword_list(value):
    if isinstance(value, list):
        parts = value
    else:
        parts = re.split(r"[,，、;\s]+", str(value or ""))
    return [str(part).strip() for part in parts if str(part).strip()][:20]


def extract_rows(payload):
    text = str(payload.get("text") or payload.get("desc") or "").strip()
    lines = clean_lines(text)
    title = str(payload.get("title") or "").strip()
    image_paths = [str(path) for path in payload.get("image_paths") or []]
    keywords = keyword_list(payload.get("keywords"))

    ip = first_match(lines, r"(?:^|\b)(?:IP|作品)\s*[:：]\s*([^。；;]+)") or title
    publisher = first_match(
        lines, r"(?:出品方?|发行|制作)\s*[:：]\s*([^。；;]+)")
    series = first_match(lines, r"(?:系列|主题)\s*[:：]\s*([^。；;]+)")
    release_date = first_match(
        lines,
        r"(?:发售|开售|上线|预约|截团|开团)[^0-9]*"
        r"(\d{4}\s*[年./-]\s*\d{1,2}(?:\s*[月./-]\s*\d{1,2})?\s*日?)",
    ) or str(payload.get("published") or "").strip()
    spec_lines = [
        line for line in lines
        if re.search(r"尺寸|规格|材质|面料|大小|cm|mm|克|kg|g\b", line, re.I)
    ]
    note_lines = [
        line for line in lines
        if re.search(r"限定|隐藏|满赠|特典|赠品|盲盒|概率", line)
    ]
    base = {
        "ip": ip,
        "publisher": publisher,
        "release_date": release_date,
        "series": series,
        "items": "",
        "spec": "\n".join(spec_lines),
        "price": "",
        "images": image_paths,
        "notes": "\n".join(note_lines),
        "source_text": text,
    }

    rows = []
    for index, line in enumerate(lines):
        prices = list(PRICE_RE.finditer(line))
        if not prices:
            continue
        prefix = line[:prices[0].start()].strip(" -:：|；;，,")
        item = prefix
        if len(item) < 2 and index:
            item = lines[index - 1].strip(" -:：|；;，,")
        for match in prices:
            row = base.copy()
            row["items"] = item or title
            row["price"] = match.group(0).strip()
            rows.append(row)

    if not rows:
        rows.append(base)
    elif keywords:
        keyword_rows = []
        for line in lines:
            if any(keyword.lower() in line.lower() for keyword in keywords):
                row = base.copy()
                row["items"] = line
                keyword_rows.append(row)
        if keyword_rows:
            rows = keyword_rows
    return rows


def markdown_cell(value):
    if isinstance(value, list):
        value = "; ".join(str(item) for item in value)
    return str(value or "").replace("|", "\\|").replace("\n", "<br>")


def markdown_image_cell(job_dir, value):
    images = []
    for raw_path in value if isinstance(value, list) else []:
        path = Path(raw_path)
        try:
            relative = path.resolve().relative_to(Path(job_dir).resolve()).as_posix()
        except (OSError, ValueError):
            relative = str(raw_path)
        images.append(f"![周边图片]({relative})")
    return "<br>".join(images) if images else "\\"


def write_result_xlsx(job_dir, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "周边明细"
    headers = [label for _, label in WEB_COLUMNS]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            "" if key == "images" else row.get(key, "")
            for key, _ in WEB_COLUMNS
        ])

    border = Border(
        left=Side(style="thin", color="D9CDB4"),
        right=Side(style="thin", color="D9CDB4"),
        top=Side(style="thin", color="D9CDB4"),
        bottom=Side(style="thin", color="D9CDB4"),
    )
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=len(WEB_COLUMNS)):
        for cell in row:
            cell.alignment = alignment
            cell.border = border
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FAF4E3")

    mergeable = {"publisher", "release_date", "ip", "series", "notes"}
    for column_index, (key, _) in enumerate(WEB_COLUMNS, 1):
        if key not in mergeable:
            continue
        start = 0
        while start < len(rows):
            value = str(rows[start].get(key) or "\\").strip()
            end = start + 1
            while end < len(rows):
                if str(rows[end].get(key) or "\\").strip() != value:
                    break
                end += 1
            if end - start > 1:
                sheet.merge_cells(
                    start_row=start + 2,
                    start_column=column_index,
                    end_row=end + 1,
                    end_column=column_index,
                )
            start = end

    widths = [18, 34, 18, 18, 22, 36, 16, 22, 48]
    for column_index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = "A2"

    image_column = next(
        index for index, (key, _) in enumerate(WEB_COLUMNS, 1)
        if key == "images"
    )
    image_letter = get_column_letter(image_column)
    for row_index, row in enumerate(rows, 2):
        images = row.get("images") or []
        if not images or not Path(images[0]).is_file():
            continue
        image = XLImage(images[0])
        scale = min(1, 140 / image.width, 140 / image.height)
        image.width = max(1, round(image.width * scale))
        image.height = max(1, round(image.height * scale))
        sheet.add_image(image, f"{image_letter}{row_index}")
        sheet.row_dimensions[row_index].height = max(
            90, image.height * 0.75 + 8
        )

    xlsx_path = job_dir / "result.xlsx"
    workbook.save(xlsx_path)
    return xlsx_path


def write_result_bundle(job_dir, payload, rows):
    public_rows = []
    for row in rows:
        public_rows.append({
            key: row.get(key, []) if key == "images" else row.get(key, "")
            for key, _ in EXPORT_COLUMNS
        })
    bundle = {
        "job_id": payload["job_id"],
        "mode": payload["mode"],
        "source": {
            "platform": payload["platform"],
            "url": payload["url"],
            "title": payload["title"],
            "author": payload["author"],
            "published": payload["published"],
        },
        "rows": public_rows,
    }
    json_path = job_dir / "result.json"
    json_path.write_text(
        json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")

    csv_path = job_dir / "result.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, label in EXPORT_COLUMNS])
        for row in public_rows:
            writer.writerow([
                "; ".join(row[key]) if isinstance(row[key], list) else row[key]
                for key, _ in EXPORT_COLUMNS
            ])

    md_path = job_dir / "result.md"
    headers = [label for _, label in WEB_COLUMNS]
    md_lines = [
        "# " + str(payload['title'] or '周边识别结果').splitlines()[0].strip(),
        "",
        f"- 来源：{payload['url']}",
        f"- 平台：{payload['platform']}",
        f"- 模式：{mode_name(payload['mode'])}",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in public_rows:
        cells = [
            markdown_image_cell(job_dir, row[key])
            if key == "images" else markdown_cell(row[key])
            for key, _ in WEB_COLUMNS
        ]
        md_lines.append("| " + " | ".join(cells) + " |")
    md_path.write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    xlsx_path = write_result_xlsx(job_dir, public_rows)
    return json_path, csv_path, md_path, xlsx_path


def sniff_ext(path):
    head = path.read_bytes()[:16]
    if head[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return ".webp"
    return ".bin"


def download_images(urls, referer, outdir):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    curl = shutil.which("curl.exe") or shutil.which("curl") or "curl"
    paths = []
    failed = []
    for i, url in enumerate(urls[:40], start=1):
        tmp = outdir / f"{i:02d}.download"
        cmd = [
            curl, "-L", "--fail", "--silent", "--show-error",
            "--max-time", "60",
            "-A", "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126 Safari/537.36",
            "-e", referer or url,
            "-o", str(tmp),
            url,
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
            if proc.returncode or not tmp.exists() or tmp.stat().st_size == 0:
                failed.append(url)
                continue
            final = outdir / f"{i:02d}{sniff_ext(tmp)}"
            tmp.replace(final)
            paths.append(str(final))
        except (OSError, subprocess.TimeoutExpired):
            failed.append(url)
        finally:
            if tmp.exists():
                tmp.unlink(missing_ok=True)
    return paths, failed


def allowed_source_url(url):
    parsed = urlparse(str(url or ""))
    host = (parsed.hostname or "").lower()
    if parsed.scheme not in ("http", "https"):
        return False
    return (
        host == "xiaohongshu.com"
        or host.endswith(".xiaohongshu.com")
        or host == "weibo.com"
        or host.endswith(".weibo.com")
        or host == "weibo.cn"
        or host.endswith(".weibo.cn")
    )


def fetch_text(url, accept="text/html,application/xhtml+xml", headers=None):
    request_headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/126.0.0.0 Safari/537.36"
        ),
        "Accept": accept,
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.7",
        "Referer": "https://www.xiaohongshu.com/" if "xiaohongshu" in url else "https://weibo.com/",
    }
    request_headers.update(headers or {})
    req = Request(url, headers=request_headers)
    try:
        with urlopen(req, timeout=30) as response:
            raw = response.read(5 * 1024 * 1024 + 1)
            if len(raw) > 5 * 1024 * 1024:
                raise RuntimeError("页面内容超过 5MB，无法处理")
            final_url = response.geturl()
            if not allowed_source_url(final_url):
                raise RuntimeError("页面跳转到了不支持的域名")
            charset = response.headers.get_content_charset() or "utf-8"
            return final_url, raw.decode(charset, errors="replace")
    except HTTPError as exc:
        raise RuntimeError(f"页面请求失败：HTTP {exc.code}") from exc
    except URLError as exc:
        raise RuntimeError(f"页面无法访问：{exc.reason}") from exc


class SourceHTMLParser(HTMLParser):
    BLOCK_TAGS = {
        "article", "aside", "blockquote", "br", "dd", "div", "dl", "dt",
        "footer", "h1", "h2", "h3", "h4", "h5", "h6", "header", "li",
        "main", "nav", "p", "section", "td", "th", "tr",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.meta = {}
        self.title = ""
        self.text = []
        self.images = []
        self.scripts = []
        self._script = None
        self._title = False

    def handle_starttag(self, tag, attrs):
        values = {key.lower(): value or "" for key, value in attrs}
        if tag == "meta":
            key = (values.get("property") or values.get("name") or "").lower()
            if key and values.get("content"):
                self.meta[key] = values["content"].strip()
        elif tag == "img":
            source = (
                values.get("src") or values.get("data-src")
                or values.get("data-original") or values.get("data-lazy-src")
            )
            if source:
                self.images.append(source.strip())
        elif tag == "script":
            self._script = []
        elif tag == "title":
            self._title = True
        if tag in self.BLOCK_TAGS:
            self.text.append("\n")

    def handle_endtag(self, tag):
        if tag == "script" and self._script is not None:
            self.scripts.append("".join(self._script))
            self._script = None
        elif tag == "title":
            self._title = False
        elif tag in self.BLOCK_TAGS:
            self.text.append("\n")

    def handle_data(self, data):
        if self._script is not None:
            self._script.append(data)
        elif self._title:
            self.title += data
        else:
            self.text.append(data)


def json_object_at(text, start):
    depth = 0
    in_string = False
    escaped = False
    for index in range(start, len(text)):
        char = text[index]
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return text[start:index + 1]
    return ""


def embedded_json(scripts):
    objects = []
    markers = ("__INITIAL_STATE__", "__NEXT_DATA__")
    for script in scripts:
        raw = script.strip()
        if not raw:
            continue
        if "application/json" in raw[:80]:
            raw = raw.split(">", 1)[-1]
        candidates = []
        if raw.startswith("{"):
            candidates.append(raw)
        for marker in markers:
            pos = raw.find(marker)
            if pos >= 0:
                brace = raw.find("{", pos)
                if brace >= 0:
                    candidates.append(json_object_at(raw, brace))
        for candidate in candidates:
            if not candidate:
                continue
            try:
                # XHS embeds JavaScript undefined in its otherwise-JSON state.
                objects.append(json.loads(
                    re.sub(r"\bundefined\b", "null", candidate)))
            except json.JSONDecodeError:
                pass
    return objects


def dict_score(value):
    keys = {
        "desc", "description", "content", "text", "title", "imageList",
        "image_list", "images", "pics", "user", "created_at", "time",
    }
    return sum(1 for key in value if key in keys)


def best_content_dict(value):
    best = None
    best_score = 0
    stack = [value]
    while stack:
        item = stack.pop()
        if isinstance(item, dict):
            score = dict_score(item)
            if score > best_score:
                best = item
                best_score = score
            stack.extend(item.values())
        elif isinstance(item, list):
            stack.extend(item)
    return best or {}


def nested_text(value):
    if isinstance(value, str):
        return unescape(value)
    if isinstance(value, dict):
        for key in ("text", "content", "desc", "description", "value"):
            if value.get(key):
                return nested_text(value[key])
    return ""


def collect_image_values(value, images):
    if isinstance(value, str):
        if value.startswith(("http://", "https://")):
            images.append(value)
    elif isinstance(value, list):
        for item in value:
            collect_image_values(item, images)
    elif isinstance(value, dict):
        for key in ("urlDefault", "urlPre", "url", "large", "original", "origin"):
            if value.get(key):
                collect_image_values(value[key], images)
                return
        for item in value.values():
            collect_image_values(item, images)


def state_content(objects):
    content = {}
    for obj in objects:
        item = best_content_dict(obj)
        if not item:
            continue
        text = nested_text(item.get("desc") or item.get("description")
                           or item.get("content") or item.get("text"))
        title = nested_text(item.get("title"))
        user = item.get("user") if isinstance(item.get("user"), dict) else {}
        author = nested_text(
            item.get("nickname") or user.get("nickname")
            or user.get("screen_name") or item.get("author")
        )
        published = nested_text(
            item.get("created_at") or item.get("createTime") or item.get("time")
        )
        images = []
        for key in ("imageList", "image_list", "images", "pics"):
            if key in item:
                collect_image_values(item[key], images)
        if text or title or images:
            content = {
                "text": text,
                "title": title,
                "author": author,
                "published": published,
                "images": images,
            }
            if text and images:
                break
    return content


def html_text(raw):
    parser = SourceHTMLParser()
    parser.feed(raw or "")
    return "\n".join(clean_lines("".join(parser.text)))


def weibo_status_id(url):
    parsed = urlparse(url)
    query_id = parse_qs(parsed.query).get("id", [""])[0]
    if re.fullmatch(r"[0-9A-Za-z]+", query_id or ""):
        return query_id
    patterns = (
        r"/status/([0-9A-Za-z]+)",
        r"/detail/([0-9A-Za-z]+)",
        r"/statuses/show/([0-9A-Za-z]+)",
        r"/\d+/([0-9A-Za-z]+)",
    )
    for pattern in patterns:
        match = re.search(pattern, parsed.path)
        if match:
            return match.group(1)
    return ""


def fetch_weibo_post(url):
    status_id = weibo_status_id(url)
    if not status_id:
        return {}
    _, raw = fetch_text(
        f"https://m.weibo.cn/statuses/show?id={quote(status_id)}",
        accept="application/json,text/plain,*/*",
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Linux; Android 14; Pixel 8) "
                "AppleWebKit/537.36 Chrome/140 Mobile Safari/537.36"
            ),
            "Referer": "https://m.weibo.cn/",
            "X-Requested-With": "XMLHttpRequest",
        },
    )
    try:
        data = json.loads(raw).get("data") or {}
    except json.JSONDecodeError:
        return {}
    if not data:
        return {}
    images = []
    for item in data.get("pics") or []:
        if not isinstance(item, dict):
            continue
        collect_image_values(
            item.get("large") or item.get("original") or item.get("url"),
            images,
        )
    return {
        "platform": "weibo",
        "title": html_text(data.get("text")) or "微博周边宣传",
        "text": html_text(data.get("text")),
        "author": nested_text((data.get("user") or {}).get("screen_name")),
        "published": nested_text(data.get("created_at")),
        "images": images,
    }


def normalize_images(urls, base_url):
    seen = set()
    result = []
    for value in urls:
        raw = str(value or "").strip()
        if not raw:
            continue
        image_url = urljoin(base_url, raw)
        if not image_url.startswith(("http://", "https://")):
            continue
        image_host = (urlparse(image_url).hostname or "").lower()
        if not any(
            image_host == host or image_host.endswith("." + host)
            for host in ("xhscdn.com", "xiaohongshu.com", "sinaimg.cn",
                         "weibo.com", "weibo.cn", "qpic.cn")
        ):
            continue
        if IMAGE_BAD_RE.search(image_url) or IMAGE_AVATAR_RE.search(image_url):
            continue
        parsed_image = urlparse(image_url)
        image_key = (
            (parsed_image.hostname or "").lower(),
            parsed_image.path,
        )
        if image_key in seen:
            continue
        seen.add(image_key)
        result.append(image_url)
        if len(result) >= 40:
            break
    return result


def parse_json_message(raw):
    text = str(raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I)
    decoder = json.JSONDecoder()
    for index, char in enumerate(text):
        if char not in "[{":
            continue
        try:
            value, _ = decoder.raw_decode(text[index:])
            return value
        except json.JSONDecodeError:
            continue
    raise RuntimeError("图片识别没有返回有效 JSON")


def refine_content_bounds(image, bounds, min_top=0):
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return bounds

    width, height = image.size
    left, top, right, bottom = bounds
    analysis_width = min(500, width)
    scale = analysis_width / width
    analysis_height = max(1, round(height * scale))
    sample = image.resize(
        (analysis_width, analysis_height), Image.Resampling.NEAREST
    )
    difference = ImageChops.difference(
        sample, Image.new("RGB", sample.size, "white")
    ).convert("L")
    mask = difference.point(lambda value: 255 if value > 18 else 0)

    def content_runs(values, threshold, gap):
        runs = []
        start = None
        last = None
        for index, value in enumerate(values):
            if value <= threshold:
                continue
            if start is None:
                start = index
            elif index - last > gap:
                runs.append((start, last))
                start = index
            last = index
        if start is not None:
            runs.append((start, last))
        return runs

    search_left = max(0, floor(left * scale))
    search_right = min(analysis_width, ceil(right * scale))
    row_data = mask.tobytes()
    row_counts = [
        row_data[
            index * analysis_width + search_left:
            index * analysis_width + search_right
        ].count(255)
        for index in range(analysis_height)
    ]
    row_runs = content_runs(
        row_counts,
        max(1, round((search_right - search_left) * 0.01)),
        max(1, round(analysis_height * 0.006)),
    )
    if not row_runs:
        return bounds

    minimum_height = max(2, round((bottom - top) * scale * 0.25))
    tall_runs = [
        run for run in row_runs
        if run[1] - run[0] + 1 >= minimum_height
    ]
    if not tall_runs:
        return bounds

    target_top = top * scale
    target_bottom = bottom * scale
    min_top_sample = floor(min_top * scale)
    if min_top:
        candidates = [
            run for run in tall_runs if run[1] + 1 > min_top_sample
        ]
        if not candidates:
            return bounds
        row_run = min(candidates, key=lambda run: run[0])
    else:
        overlapping_runs = [
            run for run in tall_runs
            if min(run[1] + 1, target_bottom) > max(run[0], target_top)
        ]
        if overlapping_runs:
            row_run = max(
                overlapping_runs,
                key=lambda run: min(run[1] + 1, target_bottom)
                - max(run[0], target_top),
            )
        else:
            target_center = (target_top + target_bottom) / 2
            row_run = min(
                tall_runs,
                key=lambda run: abs(
                    (run[0] + run[1] + 1) / 2 - target_center
                ),
            )

    refined_top = round(row_run[0] / scale)
    refined_bottom = round((row_run[1] + 1) / scale)
    if refined_top < min_top:
        refined_top = min_top
    if refined_bottom <= refined_top:
        return bounds

    band_top = max(0, floor(refined_top * scale))
    band_bottom = min(analysis_height, ceil(refined_bottom * scale))
    band = mask.crop((0, band_top, analysis_width, band_bottom))
    column_data = band.transpose(Image.Transpose.ROTATE_270).tobytes()
    band_height = band.height
    column_counts = [
        column_data[
            index * band_height:(index + 1) * band_height
        ].count(255)
        for index in range(analysis_width)
    ]
    column_runs = content_runs(
        column_counts,
        max(1, round(band_height * 0.01)),
        max(2, round(analysis_width * 0.04)),
    )
    target_left = left * scale
    target_right = right * scale
    matching_runs = [
        run for run in column_runs
        if run[1] + 1 >= target_left and run[0] <= target_right
    ]
    if not matching_runs:
        return (left, refined_top, right, refined_bottom)

    refined_left = round(min(run[0] for run in matching_runs) / scale)
    refined_right = round(
        (max(run[1] for run in matching_runs) + 1) / scale
    )
    return (refined_left, refined_top, refined_right, refined_bottom)


def crop_image_region(source_path, box, output_path, min_top=0, refine=True):
    try:
        from PIL import Image, ImageChops, ImageOps
    except ImportError:
        return False

    try:
        values = [float(value) for value in box]
    except (TypeError, ValueError):
        return False
    if len(values) != 4:
        return False

    with Image.open(source_path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        width, height = image.size
        if max(values) > 1000:
            left, top, right, bottom = values
        elif max(values) <= 1:
            left, top, right, bottom = (
                values[0] * width,
                values[1] * height,
                values[2] * width,
                values[3] * height,
            )
        else:
            left, top, right, bottom = (
                values[0] / 1000 * width,
                values[1] / 1000 * height,
                values[2] / 1000 * width,
                values[3] / 1000 * height,
            )

        content_bounds = (
            refine_content_bounds(
                image, (round(left), round(top), round(right), round(bottom)),
                min_top,
            )
            if refine
            else (round(left), round(top), round(right), round(bottom))
        )
        left, top, right, bottom = content_bounds
        crop_width = max(1, right - left)
        crop_height = max(1, bottom - top)
        pad_x = max(3, round(crop_width * 0.01))
        pad_y = max(3, round(crop_height * 0.01))
        bounds = (
            max(0, left - pad_x),
            max(0, top - pad_y),
            min(width, right + pad_x),
            min(height, bottom + pad_y),
        )
        if bounds[2] - bounds[0] < 20 or bounds[3] - bounds[1] < 20:
            return False

        cropped = image.crop(bounds)
        background = cropped.getpixel((0, 0))
        difference = ImageChops.difference(
            cropped, Image.new("RGB", cropped.size, background)
        ).convert("L")
        content_box = difference.point(
            lambda value: 255 if value > 18 else 0
        ).getbbox()
        if content_box:
            trimmed = cropped.crop(content_box)
            if trimmed.width >= 20 and trimmed.height >= 20:
                cropped = trimmed

        output_path.parent.mkdir(parents=True, exist_ok=True)
        cropped.save(output_path, quality=95, optimize=True)
        return content_bounds
    return False


def move_unit_sale_to_price(row):
    notes = row.get("notes", "").strip()
    match = UNIT_SALE_RE.match(notes)
    if not match:
        return
    unit_sale = match.group(1)
    price = row.get("price", "").strip()
    row["price"] = f"{price}，{unit_sale}" if price else unit_sale
    row["notes"] = notes[match.end():].strip()


def prepare_vision_images(image_paths, job_dir):
    parts = []
    try:
        from PIL import Image, ImageOps
    except ImportError:
        return [
            {
                "path": Path(path),
                "source_path": Path(path),
                "source_height": 0,
                "slice_top": 0,
                "slice_height": 0,
            }
            for path in image_paths
        ]

    images = []
    for source_index, raw_path in enumerate(image_paths, 1):
        path = Path(raw_path)
        try:
            with Image.open(path) as source:
                image = ImageOps.exif_transpose(source)
                width, height = image.size
                images.append({
                    "source_index": source_index,
                    "path": path,
                    "width": width,
                    "height": height,
                    "image": image.convert("RGB"),
                })
        except (OSError, ValueError):
            continue

    long_images = [
        item for item in images if item["height"] > VISION_SLICE_HEIGHT
    ]
    # ponytail: only drop thumbnail-like extras when several full catalog
    # images are present; a single long image may coexist with unique photos.
    long_indexes = {item["source_index"] for item in long_images}
    if len(long_images) >= 2 and all(
        item["height"] <= 1200 and item["width"] <= 1200
        for item in images if item["source_index"] not in long_indexes
    ):
        selected_images = long_images
    else:
        selected_images = images

    slices_dir = job_dir / "images" / "vision_slices"
    for item in selected_images:
        if item["height"] <= VISION_SLICE_HEIGHT:
            parts.append({
                "path": item["path"],
                "source_path": item["path"],
                "source_height": item["height"],
                "slice_top": 0,
                "slice_height": item["height"],
            })
            continue

        slices_dir.mkdir(parents=True, exist_ok=True)
        for slice_number, top in enumerate(
            range(0, item["height"], VISION_SLICE_HEIGHT), 1
        ):
            bottom = min(item["height"], top + VISION_SLICE_HEIGHT)
            slice_path = slices_dir / (
                f"{item['source_index']:02d}_{slice_number:02d}"
                f"_{top}-{bottom}.jpg"
            )
            item["image"].crop((0, top, item["width"], bottom)).save(
                slice_path, quality=95, optimize=True
            )
            parts.append({
                "path": slice_path,
                "source_path": item["path"],
                "source_height": item["height"],
                "slice_top": top,
                "slice_height": bottom - top,
            })
    return parts


def region_on_original(part, box):
    try:
        values = [float(value) for value in box]
    except (TypeError, ValueError):
        return None
    if len(values) != 4 or not part.get("source_height"):
        return None
    if max(values) > 1000:
        left, top, right, bottom = values
    elif max(values) <= 1:
        left, top, right, bottom = (value * 1000 for value in values)
    else:
        left, top, right, bottom = values

    slice_top = part["slice_top"]
    slice_height = part["slice_height"]
    source_height = part["source_height"]
    original_top = slice_top + top / 1000 * slice_height
    original_bottom = slice_top + bottom / 1000 * slice_height
    return [
        left,
        original_top / source_height * 1000,
        right,
        original_bottom / source_height * 1000,
    ]


def merge_regions_by_source(regions, vision_parts):
    grouped = {}
    for region in regions:
        if not isinstance(region, dict):
            continue
        try:
            index = int(region.get("image_index"))
        except (TypeError, ValueError):
            continue
        if not 1 <= index <= len(vision_parts):
            continue
        part = vision_parts[index - 1]
        box = region_on_original(part, region.get("box"))
        if not box:
            continue
        source_key = str(part["source_path"])
        if source_key not in grouped:
            grouped[source_key] = [part["source_path"], box]
            continue
        merged = grouped[source_key][1]
        grouped[source_key][1] = [
            min(merged[0], box[0]),
            min(merged[1], box[1]),
            max(merged[2], box[2]),
            max(merged[3], box[3]),
        ]
    return list(grouped.values())


def normalize_series_from_spec(rows, row_sources):
    series_counts = {}
    for row, sources in zip(rows, row_sources):
        series = row.get("series", "").strip()
        if not series or series == "\\":
            continue
        for source in sources:
            counts = series_counts.setdefault(source, {})
            counts[series] = counts.get(series, 0) + 1

    for row, sources in zip(rows, row_sources):
        series = row.get("series", "").strip()
        if series not in ("", "\\") and series != row.get("ip", "").strip():
            continue
        label_text = " ".join(re.findall(r"【([^】]+)】", row.get("spec", "")))
        votes = {}
        for source in sources:
            counts = series_counts.get(source, {})
            for candidate, count in counts.items():
                if candidate and candidate in label_text:
                    votes[candidate] = votes.get(candidate, 0) + count
        if votes:
            row["series"] = max(votes, key=votes.get)


def vision_rows(cfg, payload, image_paths):
    cli = resolve_codex_cli(cfg)
    if not cli:
        raise RuntimeError("未找到 codex 命令，无法进行图片识别")
    if not image_paths:
        raise RuntimeError("没有可用于图片识别的商品图")

    job_dir = Path(payload["job_dir"])
    response_path = job_dir / "vision-response.txt"
    response_path.unlink(missing_ok=True)
    keywords = "、".join(payload.get("keywords") or []) or "无"
    vision_parts = prepare_vision_images(image_paths, job_dir)
    if not vision_parts:
        raise RuntimeError("没有可用于图片识别的商品图")
    prompt = f"""你是周边商品明细提取器。直接目测图片并立即输出 JSON；不要执行命令、调用工具或分析像素。只依据下面的帖文和附图，不要访问网页，不要猜测，也不要补全看不清的内容。

帖文标题：{payload.get('title') or ''}
帖文网址：{payload.get('url') or ''}
筛选关键词：{keywords}
帖文正文：
{str(payload.get('text') or '')[:12000]}

请逐张检查全部附图。为保留长图中的小字，部分长图已按从上到下的顺序切成多个片段；每个片段都是独立附图，并按下文 image_index 对应。请遵守：
1. 一张图片可能并列展示多个商品套组；每个编号、标题或独立套组各输出一行。前后商品块排版相似也不能合并或略过，必须逐块核对。
2. 不同图片片段即使内容结构相似，也必须分别识别，不能因为结构相同而合并或略过。
3. 满赠纸袋或特典不单独增加贩售商品行，将其名称、尺寸、满赠条件和“不可叠加”等信息写入每条相关记录的 notes；不要依据商品图中的局部款式说明生成 notes。
4. publisher 读取“出品/出品方”，不要把“原著、授权、承制”或社交账号当作出品方；同时出现“授权”和“出品”时必须取“出品”后的主体。
5. series 从商品块上方或附近最近的“系列/主题”标题读取，去掉方括号、书名号和结尾的“系列”字样，例如“【Golden Hour】系列”应输出“Golden Hour”；同一商品图或切片主系列下的所有商品通常使用同一个 series，不要因 spec 内同时出现多个系列名而改写 series。不要附加 1、2、3 等编号，只有确实没有系列信息时才写“\\”。
6. release_date 综合正文和图片，按“发售日期丨地点丨具体展位”整理。
7. spec 必须优先读取商品图下方和旁边的小字。只要尺寸、材质、工艺中任意一项存在，就必须把已读取到的项全部写入；只对确实缺失的子项写“\\”，严禁因为没有同时找到三项而整项留空。多个制品用“；”分隔，格式如“【徽章】尺寸约80x80mm 材质马口铁、PET 工艺细沙镭射底、烫哑金”。
8. price 只记录售价、计价单位和销售数量或方式，例如“42CNY/组”“69CNY/只，2只/套，按套售卖”，不要把满赠、尺寸等信息写入 price。
9. notes 汇总正文和图片中的所有 72H 特典及各级满赠，完整记录赠送对象、所属系列、数量、尺寸、门槛和限制；72H 是活动时限，绝不能误写成“满72CNY”等价格。同一批商品的共同备注必须完全相同，不要把单个商品的售价或销售数量写入 notes。
10. image_regions 是该行商品图在当前附图中的紧边界数组。image_index 是从 1 开始的附图编号；box 使用 [左, 上, 右, 下] 归一化坐标，范围 0-1000。
11. 只框商品照片、套组图或商品效果图，严格排除编号、标题、规格、价格、分隔线、背景、无关文字和相邻商品；不要按固定高度切块，也不要返回整行或整张图片。默认每行只用多个商品图的合并紧边界；只有商品图相隔很远或属于明显不同套组时才返回多个框。
12. 所有无法确认的字段使用“\\”，不要编造。

只输出以下结构的 JSON，不要 Markdown，不要解释：
{{"rows":[{{"publisher":"","release_date":"","series":"","ip":"","items":"","spec":"","price":"","image_regions":[{{"image_index":1,"box":[420,575,920,685]}}],"notes":""}}]}}
"""
    cmd = [
        cli, "exec", "--ephemeral", "--skip-git-repo-check",
        "--sandbox", "read-only", "--color", "never",
        "-c", "mcp_servers.obsidian.enabled=false",
        "-c", "mcp_servers.baidu_netdisk.enabled=false",
        "-c", "mcp_servers.cn-scraper.enabled=false",
        "-c", "mcp_servers.node_repl.enabled=false",
        "-o", str(response_path), "-i",
    ]
    cmd.extend(str(part["path"]) for part in vision_parts)
    cmd.extend(["--", prompt])
    try:
        proc = subprocess.run(
            cmd, cwd=str(ROOT), capture_output=True, text=True,
            encoding="utf-8", errors="replace", stdin=subprocess.DEVNULL,
            timeout=360,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("图片识别超时") from exc
    if not response_path.exists():
        detail = " ".join((proc.stderr or proc.stdout or "").split())[:240]
        raise RuntimeError(
            f"图片识别失败：{detail or f'进程退出码 {proc.returncode}'}"
        )

    try:
        parsed = parse_json_message(response_path.read_text(
            encoding="utf-8", errors="replace"))
    except RuntimeError:
        detail = " ".join((proc.stderr or proc.stdout or "").split())[:240]
        raise RuntimeError(
            f"图片识别失败：{detail or '未返回有效结果'}"
        ) from None
    items = parsed.get("rows") if isinstance(parsed, dict) else parsed
    if not isinstance(items, list):
        raise RuntimeError("图片识别结果缺少 rows")

    aliases = {
        "出品方": "publisher",
        "发售时间及地点": "release_date",
        "系列": "series",
        "IP名称": "ip",
        "周边明细": "items",
        "尺寸丨材质丨工艺": "spec",
        "价格": "price",
        "图片": "image_regions",
        "备注": "notes",
    }
    rows = []
    row_sources = []
    crops_dir = job_dir / "images" / "crops"
    for row_number, item in enumerate(items, 1):
        if not isinstance(item, dict):
            continue
        item = {aliases.get(key, key): value for key, value in item.items()}
        row = {
            key: str(item.get(key) or "").strip()
            for key in (
                "publisher", "release_date", "series", "ip", "items", "spec",
                "price", "notes",
            )
        }
        move_unit_sale_to_price(row)
        row["images"] = []
        regions = item.get("image_regions")
        if not isinstance(regions, list):
            regions = []
        sources = set()
        for region_number, (source_path, original_box) in enumerate(
            merge_regions_by_source(regions, vision_parts), 1
        ):
            sources.add(str(source_path))
            crop_path = crops_dir / (
                f"{row_number:02d}_{region_number:02d}_{source_path.stem}.jpg"
            )
            bounds = crop_image_region(
                source_path, original_box, crop_path, refine=False,
            )
            if bounds:
                row["images"].append(str(crop_path))

        row["source_text"] = str(payload.get("text") or "")
        rows.append(row)
        row_sources.append(sources)
    normalize_series_from_spec(rows, row_sources)
    return rows


def analyze_source(url, keywords):
    if not allowed_source_url(url):
        raise RuntimeError("只支持小红书或微博网址")
    platform = "xiaohongshu" if "xiaohongshu" in urlparse(url).hostname else "weibo"
    if platform == "weibo":
        try:
            source = fetch_weibo_post(url)
            if source:
                source["images"] = normalize_images(source.get("images") or [], url)
                source["url"] = url
                return source
        except RuntimeError:
            pass

    final_url, raw = fetch_text(url)
    parser = SourceHTMLParser()
    parser.feed(raw)
    state = state_content(embedded_json(parser.scripts))
    text = state.get("text") or "\n".join(clean_lines("".join(parser.text)))
    title = (
        state.get("title") or parser.meta.get("og:title")
        or parser.title or "周边宣传"
    )
    author = (
        state.get("author") or parser.meta.get("og:author")
        or parser.meta.get("author") or ""
    )
    published = (
        state.get("published")
        or parser.meta.get("article:published_time") or ""
    )
    images = normalize_images(
        (state.get("images") or []) + [parser.meta.get("og:image", "")]
        + parser.images,
        final_url,
    )
    if not text and not images:
        raise RuntimeError("未读取到正文或图片，请确认帖子可公开访问")
    return {
        "platform": platform,
        "url": final_url,
        "title": title.strip(),
        "author": author.strip(),
        "published": published.strip(),
        "text": text.strip()[:50000],
        "images": images,
    }


def output_url(path):
    if not path:
        return ""
    try:
        relative = Path(path).resolve().relative_to(ROOT).as_posix()
    except (OSError, ValueError):
        return ""
    return "/" + quote(relative)


def web_analyze(data):
    url = str(data.get("url") or "").strip()
    keywords = keyword_list(data.get("keywords"))
    source = analyze_source(url, keywords)
    payload = {
        "mode": "local",
        "platform": source["platform"],
        "url": source["url"],
        "title": source["title"],
        "author": source["author"],
        "published": source["published"],
        "text": source["text"],
        "images": source["images"],
        "keywords": keywords,
        "use_vision": True,
    }
    job = archive_job(payload)
    if job.get("vision_error"):
        raise RuntimeError("图片识别失败：" + job["vision_error"])
    bundle = json.loads(Path(job["result_json"]).read_text(encoding="utf-8"))
    rows = []
    for row in bundle.get("rows") or []:
        item = {}
        for key, _ in WEB_COLUMNS:
            value = row.get(key) or ""
            if key == "images":
                links = []
                for path in value:
                    try:
                        relative = Path(path).resolve().relative_to(ROOT).as_posix()
                    except (OSError, ValueError):
                        continue
                    links.append("/" + quote(relative))
                item[key] = links
            else:
                item[key] = value
        rows.append(item)
    return {
        "ok": True,
        "job_id": job["job_id"],
        "source": {
            "platform": source["platform"],
            "url": source["url"],
            "title": source["title"],
        },
        "rows": rows,
        "md_url": output_url(job.get("result_md")),
        "xlsx_url": output_url(job.get("result_xlsx")),
        "notice": job.get("vision_error") or "",
    }


def queue_to_codex(cfg, source_path, job_dir, data):
    cli = resolve_codex_cli(cfg)
    thread = cfg.get("codex_thread") or ""
    if not cli or not Path(cli).exists():
        raise RuntimeError("未找到 codex 命令，请配置 config.json 的 codex_cli")
    if not thread:
        raise RuntimeError("config.json 未配置 codex_thread（当前 Codex 会话 ID）")

    result_file = job_dir / "result.txt"
    target = target_name(data.get("archive_target"))
    archive_dir = str(Path(data.get("archive_dir") or "").expanduser())
    msg = (
        f"浏览器插件提交了一页周边宣传。源数据在 {source_path}，"
        f"图片已下载到 {job_dir / 'images'}，统一结果包在 {job_dir / 'result.json'}。\n"
        "请读取正文和图片，识别官方宣传中的周边信息，并直接完成归档。\n"
        "先补全或修正 result.json 的 rows，再同步更新 result.csv 和 result.md。\n"
        "识别字段：IP、出品方、上线时间、系列名、制品明细、材质规格、"
        "价格（写成 X元/个 或 X元/对 等原文形式）、图片、备注。"
        "备注只填写限定、隐藏、满赠内容，没有就留空。\n"
        "这是通用周边识别：不要确认或筛选是否与琵琶相关，也不要因为题材不是琵琶而跳过。\n"
        f"归档目标：{target}；归档位置：{archive_dir}。"
        "按目标软件可用格式写入：Obsidian 用 Markdown 并保存图片附件；"
        "WPS 用表格或 CSV；印象笔记用可导入的 ENEX 或 HTML；其他软件选择合适格式。\n"
        f"完成后把归档文件的绝对路径写入 {result_file}，并简要汇报结果。"
    )
    proc = subprocess.run(
        [cli, "queue", "--thread", thread, "--message", msg],
        cwd=str(ROOT), capture_output=True, text=True, timeout=60)
    if proc.returncode:
        detail = (proc.stderr or proc.stdout or "").strip()[:500]
        raise RuntimeError(f"Codex queue 失败：{detail}")


def archive_job(data):
    cfg = load_config()
    mode = str(data.get("mode") or "ai").strip().lower()
    if mode not in MODES:
        raise RuntimeError("未知的处理方式")
    text = (data.get("text") or data.get("desc") or "").strip()
    images = data.get("images") or []
    target = str(data.get("archive_target") or "").strip()
    archive_dir = str(data.get("archive_dir") or "").strip()
    if not text and not images:
        raise RuntimeError("没有正文也没有图片，无法归档")
    if mode == "ai":
        if not target:
            raise RuntimeError("请在插件里选择归档软件")
        if not archive_dir:
            raise RuntimeError("请在插件里填写归档位置")

    job_id = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + os.urandom(3).hex()
    base_dir = Path(
        (archive_dir if mode == "local" else "") or cfg.get("output_dir")
        or (ROOT / "output")
    ).expanduser()
    if not base_dir.is_absolute():
        base_dir = ROOT / base_dir
    job_dir = base_dir / job_id
    image_dir = job_dir / "images"
    job_dir.mkdir(parents=True, exist_ok=True)

    image_paths, failed = download_images(images, data.get("url") or "", image_dir)
    payload = {
        "job_id": job_id,
        "mode": mode,
        "collected_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source_type": "zhoubian-edge-collector",
        "platform": data.get("platform") or "unknown",
        "url": data.get("url") or "",
        "title": data.get("title") or "",
        "author": data.get("author") or "",
        "published": data.get("published") or "",
        "text": text,
        "desc": data.get("desc") or "",
        "image_urls": images,
        "image_paths": image_paths,
        "download_failed": failed,
        "archive_target": target,
        "archive_dir": archive_dir,
        "keywords": keyword_list(data.get("keywords")),
    }
    source_path = job_dir / "source.json"
    source_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    rows = extract_rows(payload)
    vision_error = ""
    vision_used = False
    if mode == "local" and data.get("use_vision"):
        payload["job_dir"] = str(job_dir)
        try:
            visual_rows = vision_rows(cfg, payload, image_paths)
            if visual_rows:
                rows = visual_rows
                vision_used = True
            else:
                raise RuntimeError("图片识别没有生成明细")
        except RuntimeError as exc:
            vision_error = str(exc)
    json_path, csv_path, md_path, xlsx_path = write_result_bundle(
        job_dir, payload, rows)

    if mode == "ai":
        queue_to_codex(cfg, source_path, job_dir, payload)
        note = f"；{len(failed)} 张图下载失败，Codex 会用原始链接补" if failed else ""
        message = (
            f"已交给 Codex（任务 {job_id}，图片 {len(image_paths)} 张，"
            f"归档到 {target_name(target)}）{note}"
        ).rstrip()
    else:
        note = f"；{len(failed)} 张图下载失败" if failed else ""
        if vision_error:
            note += f"；图片识别失败，已回退本地解析：{vision_error}"
        message = (
            f"已生成 {len(rows)} 行"
            f"{'图片识别' if vision_used else '本地'}结果（任务 {job_id}，"
            f"图片 {len(image_paths)} 张{note}）。结果：{job_dir}"
        )
    return {
        "ok": True,
        "job_id": job_id,
        "mode": mode,
        "job_dir": str(job_dir),
        "result_json": str(json_path),
        "result_csv": str(csv_path),
        "result_md": str(md_path),
        "result_xlsx": str(xlsx_path or ""),
        "images": len(image_paths),
        "message": message,
        "vision_used": vision_used,
        "vision_error": vision_error,
    }


class Handler(BaseHTTPRequestHandler):
    server_version = "ZhoubianBridge/0.3"

    def _allowed_origin(self):
        origin = self.headers.get("Origin") or ""
        parsed = urlparse(origin)
        return (
            not origin
            or origin.startswith("chrome-extension://")
            or origin.startswith("moz-extension://")
            or (
                parsed.scheme in ("http", "https")
                and parsed.hostname in ("127.0.0.1", "localhost")
            )
        )

    def _cors_headers(self):
        origin = self.headers.get("Origin") or ""
        if origin:
            self.send_header("Access-Control-Allow-Origin", origin)
        else:
            self.send_header("Access-Control-Allow-Origin", "*")

    def _json(self, status, obj):
        raw = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self._cors_headers()
        self.end_headers()
        self.wfile.write(raw)

    def _file(self, path):
        try:
            raw = Path(path).read_bytes()
        except OSError:
            self._json(404, {"ok": False, "error": "not found"})
            return
        content_type = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
        if content_type.startswith("text/") or content_type in (
            "application/javascript", "application/json",
        ):
            content_type += "; charset=utf-8"
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(raw)

    def _output_file(self, request_path):
        cfg = load_config()
        output_dir = Path(cfg.get("output_dir") or "output").expanduser()
        if not output_dir.is_absolute():
            output_dir = ROOT / output_dir
        relative = unquote(request_path.removeprefix("/output/"))
        target = (output_dir / relative).resolve()
        try:
            target.relative_to(output_dir.resolve())
        except ValueError:
            self._json(403, {"ok": False, "error": "forbidden"})
            return
        self._file(target)

    def do_OPTIONS(self):
        if not self._allowed_origin():
            self._json(403, {"ok": False, "error": "origin not allowed"})
            return
        self.send_response(204)
        self._cors_headers()
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self):
        if not self._allowed_origin():
            self._json(403, {"ok": False, "error": "origin not allowed"})
            return
        path = urlparse(self.path).path
        if path in ("/", "/index.html"):
            self._file(ROOT / "index.html")
        elif path == "/style.css":
            self._file(ROOT / "style.css")
        elif path == "/app.js":
            self._file(ROOT / "app.js")
        elif path == "/health":
            cfg = load_config()
            self._json(200, {
                "ok": True,
                "service": "zhoubian-edge-collector bridge",
                "codex_thread": bool(cfg.get("codex_thread")),
                "codex_cli": bool(resolve_codex_cli(cfg)),
            })
        elif path.startswith("/output/"):
            self._output_file(path)
        elif path == "/favicon.ico":
            self.send_response(204)
            self.end_headers()
        else:
            self._json(404, {"ok": False, "error": "not found"})

    def do_POST(self):
        try:
            if not self._allowed_origin():
                raise PermissionError("请求来源不受信任")
            length = int(self.headers.get("Content-Length") or 0)
            if length <= 0 or length > 3 * 1024 * 1024:
                raise RuntimeError("请求体为空或超过 3MB")
            data = json.loads(self.rfile.read(length).decode("utf-8"))
            if self.path == "/api/queue":
                self._json(200, archive_job(data))
            elif self.path == "/api/analyze":
                self._json(200, web_analyze(data))
            else:
                self._json(404, {"ok": False, "error": "not found"})
        except json.JSONDecodeError as exc:
            self._json(400, {"ok": False, "error": "请求不是有效 JSON：" + str(exc)})
        except PermissionError as exc:
            self._json(403, {"ok": False, "error": str(exc)})
        except RuntimeError as exc:
            self._json(400, {"ok": False, "error": str(exc)})
        except Exception as exc:
            self._json(500, {"ok": False, "error": str(exc)})

    def log_message(self, fmt, *args):
        sys.stdout.write(f"[zhoubian-bridge] {self.address_string()} - {fmt % args}\n")


def self_test():
    with tempfile.TemporaryDirectory() as tmp:
        sample = Path(tmp) / "a.download"
        sample.write_bytes(b"\xff\xd8\xff\xe0" + b"0" * 20)
        assert sniff_ext(sample) == ".jpg"
        sample.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 20)
        assert sniff_ext(sample) == ".png"
        payload = {
            "job_id": "test",
            "mode": "local",
            "platform": "weibo",
            "url": "https://example.com/post",
            "title": "测试系列 周边",
            "author": "",
            "published": "",
            "text": (
                "IP：测试作品\n"
                "系列：测试系列\n"
                "亚克力立牌\n"
                "尺寸：10cm\n"
                "材质：亚克力\n"
                "价格：39元/个\n"
                "满赠：购买套装赠送贴纸"
            ),
            "image_paths": [str(Path(tmp) / "01.jpg")],
        }
        rows = extract_rows(payload)
        assert len(rows) == 1
        assert rows[0]["ip"] == "测试作品"
        assert rows[0]["series"] == "测试系列"
        assert rows[0]["price"] == "39元/个"
        assert "尺寸：10cm" in rows[0]["spec"]
        assert "满赠" in rows[0]["notes"]
        price_row = {
            "price": "69CNY/只",
            "notes": "2只/套，按套售卖；满赠特典纸袋：赠品不可叠加。",
        }
        move_unit_sale_to_price(price_row)
        assert price_row["price"] == "69CNY/只，2只/套，按套售卖"
        assert price_row["notes"] == "满赠特典纸袋：赠品不可叠加。"
        mixed_series = [
            {
                "series": "测试作品",
                "ip": "测试作品",
                "spec": "【序曲海报】400x600mm；【浪漫世纪海报】400x700mm",
            },
            {
                "series": "序曲",
                "ip": "测试作品",
                "spec": "尺寸80x80mm",
            },
        ]
        normalize_series_from_spec(mixed_series, [{"11.jpg"}, {"11.jpg"}])
        assert mixed_series[0]["series"] == "序曲"
        assert mixed_series[1]["series"] == "序曲"
        json_path, csv_path, md_path, xlsx_path = write_result_bundle(
            Path(tmp), payload, rows)
        assert json.loads(json_path.read_text(encoding="utf-8"))["rows"]
        assert "39元/个" in csv_path.read_text(encoding="utf-8-sig")
        md_text = md_path.read_text(encoding="utf-8")
        assert "测试系列" in md_text
        assert "![周边图片](" in md_text
        if xlsx_path:
            assert xlsx_path.exists()
        try:
            from PIL import Image, ImageDraw
        except ImportError:
            pass
        else:
            source_path = Path(tmp) / "source.jpg"
            crop_path = Path(tmp) / "crop.jpg"
            Image.new("RGB", (100, 200), "white").save(source_path)
            assert crop_image_region(
                source_path, [0, 250, 1000, 750], crop_path)
            with Image.open(crop_path) as cropped:
                assert cropped.width == 100
                assert 100 <= cropped.height <= 120

            source_path = Path(tmp) / "product-row.jpg"
            crop_path = Path(tmp) / "product-crop.jpg"
            image = Image.new("RGB", (200, 200), "white")
            draw = ImageDraw.Draw(image)
            draw.rectangle((0, 0, 20, 199), fill="black")
            draw.rectangle((40, 80, 120, 160), fill="red")
            image.save(source_path)
            assert crop_image_region(
                source_path, [200, 400, 600, 800], crop_path)
            with Image.open(crop_path) as cropped:
                assert 80 <= cropped.width <= 84
                assert 80 <= cropped.height <= 84
                red, green, blue = cropped.getpixel((0, 0))
                assert red > 150 and red > green * 4 and red > blue * 4
            direct_path = Path(tmp) / "product-direct-crop.jpg"
            assert crop_image_region(
                source_path, [200, 400, 600, 800], direct_path,
                refine=False,
            )
            with Image.open(direct_path) as cropped:
                assert 80 <= cropped.width <= 84
                assert 80 <= cropped.height <= 84
                red, green, blue = cropped.getpixel((0, 0))
                assert red > 150 and red > green * 4 and red > blue * 4
            xlsx_test = write_result_xlsx(
                Path(tmp),
                [{"items": "测试商品", "images": [str(direct_path)]}],
            )
            if xlsx_test:
                from openpyxl import load_workbook
                assert len(load_workbook(xlsx_test).active._images) == 1

            image = Image.new("RGB", (200, 400), "white")
            draw = ImageDraw.Draw(image)
            draw.rectangle((40, 40, 120, 100), fill="red")
            draw.rectangle((40, 180, 120, 260), fill="blue")
            bounds = refine_content_bounds(
                image, (40, 150, 120, 270), min_top=140)
            assert bounds[1] >= 180

            tall_path = Path(tmp) / "tall.jpg"
            tall_path_2 = Path(tmp) / "tall-2.jpg"
            thumbnail_path = Path(tmp) / "thumb.jpg"
            Image.new("RGB", (100, 6200), "white").save(tall_path)
            Image.new("RGB", (100, 3200), "gray").save(tall_path_2)
            Image.new("RGB", (100, 100), "black").save(thumbnail_path)
            parts = prepare_vision_images(
                [str(tall_path), str(thumbnail_path), str(tall_path_2)],
                Path(tmp),
            )
            assert len(parts) == 5
            assert all(part["source_path"] != thumbnail_path for part in parts)
            assert [part["slice_top"] for part in parts[:3]] == [
                0, 3000, 6000
            ]
            assert [part["slice_height"] for part in parts[:3]] == [
                3000, 3000, 200
            ]
            mapped = region_on_original(parts[1], [100, 250, 900, 750])
            assert abs(mapped[1] - 604.84) < 0.1
            assert abs(mapped[3] - 846.77) < 0.1
            merged = merge_regions_by_source([
                {"image_index": 1, "box": [200, 800, 900, 1000]},
                {"image_index": 2, "box": [100, 0, 800, 100]},
            ], parts)
            assert len(merged) == 1
            assert merged[0][0] == tall_path
            merged_box = merged[0][1]
            assert merged_box[0] == 100 and merged_box[2] == 900
            assert abs(merged_box[1] - 387.1) < 0.1
            assert abs(merged_box[3] - 532.26) < 0.1
    assert target_name("wps") == "WPS"
    assert target_name("evernote") == "印象笔记"
    assert mode_name("local") == "无 AI 本地导出"
    assert allowed_source_url("https://www.xiaohongshu.com/explore/example")
    assert allowed_source_url("https://m.weibo.cn/detail/123")
    assert not allowed_source_url("https://example.com/post")
    assert weibo_status_id(
        "https://weibo.com/7910915063/RhhYAn0gk"
    ) == "RhhYAn0gk"
    assert weibo_status_id("https://m.weibo.cn/status/RhhYAn0gk") == (
        "RhhYAn0gk"
    )
    assert weibo_status_id("https://m.weibo.cn/detail/5341206805482556") == (
        "5341206805482556"
    )
    state = embedded_json([
        'window.__INITIAL_STATE__={"note":{"desc":"立牌 39元/个",'
        '"imageList":[{"urlDefault":"https://example.com/a.jpg"}]}}'
    ])
    parsed = state_content(state)
    assert parsed["text"] == "立牌 39元/个"
    assert parsed["images"] == ["https://example.com/a.jpg"]
    state = embedded_json([
        'window.__INITIAL_STATE__={"note":{"desc":"立牌 39元/个",'
        '"missing":undefined,"imageList":['
        '{"urlDefault":"https://example.com/a.jpg"},'
        '{"urlDefault":"https://example.com/b.jpg"}]}}'
    ])
    parsed = state_content(state)
    assert parsed["images"] == [
        "https://example.com/a.jpg", "https://example.com/b.jpg"
    ]
    assert parse_json_message(
        '```json\n{"rows":[{"publisher":"长佩文学"}]}\n```'
    )["rows"][0]["publisher"] == "长佩文学"
    assert normalize_images([
        "https://sns-webpic.xhscdn.com/a.jpg",
        "http://127.0.0.1/private.jpg",
    ], "https://www.xiaohongshu.com/") == [
        "https://sns-webpic.xhscdn.com/a.jpg"
    ]
    keyword_rows = extract_rows({
        "text": "徽章 10元/个\n亚克力立牌 39元/个",
        "title": "测试周边",
        "keywords": ["立牌"],
    })
    assert len(keyword_rows) == 1
    assert "立牌" in keyword_rows[0]["items"]
    print("self-test OK")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
        return
    cfg = load_config()
    if not cfg.get("codex_thread"):
        print(
            "config.json 缺 codex_thread，AI 模式不可用；无 AI 模式不受影响。",
            file=sys.stderr,
        )
    httpd = ThreadingHTTPServer((cfg["host"], cfg["port"]), Handler)
    print(f"Zhoubian bridge ready: http://{cfg['host']}:{cfg['port']}/")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
